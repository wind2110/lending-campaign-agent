"""Bo cot lead_id khoi sheet fact_loan (1 khach co the co nhieu lead truoc khi vay nen khong the gan 1 lead_id cho 1 khoan vay).

Thay bang lien ket qua Customer_id: voi cac lead First Loan da chuyen doi thanh khoan vay, bo sung Customer_id
vao fact_lead (khach duoc nhan dien sau khi xac thuc SDT). Lead Re-loan da co san Customer_id.
Doc va ghi de len cung 1 file (mac dinh data/full_schema_mock_v2.xlsx).
Chay: ./venv/Scripts/python.exe scripts/drop_loan_lead_id.py [duong_dan_file]
"""
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

PATH = Path(sys.argv[1] if len(sys.argv) > 1 else "data/full_schema_mock_v2.xlsx")

before = pd.read_excel(PATH, sheet_name=None)
loan0, lead0 = before["fact_loan"], before["fact_lead"]
assert "lead_id" in loan0.columns, "fact_loan da khong con cot lead_id"

wb = load_workbook(PATH)
ws_lead, ws_loan = wb["fact_lead"], wb["fact_loan"]
lead_hdr = [c.value for c in ws_lead[1]]
loan_hdr = [c.value for c in ws_loan[1]]
i_lead_id, i_lead_cust = lead_hdr.index("lead_id"), lead_hdr.index("Customer_id")
j_lead, j_cust, j_prod = loan_hdr.index("lead_id"), loan_hdr.index("Customer_id"), loan_hdr.index("product_name")

first_loan_cust = {
    row[j_lead]: row[j_cust]
    for row in ws_loan.iter_rows(min_row=2, values_only=True)
    if row[j_prod] == "First Loan" and row[j_lead]
}
n_filled = 0
for row in ws_lead.iter_rows(min_row=2):
    cust = first_loan_cust.get(row[i_lead_id].value)
    if cust and row[i_lead_cust].value in (None, ""):
        row[i_lead_cust].value = cust
        n_filled += 1
ws_loan.delete_cols(j_lead + 1)
wb.save(PATH)

# Kiem tra: doc lai file vua ghi
after = pd.read_excel(PATH, sheet_name=None)
loan1, lead1 = after["fact_loan"], after["fact_lead"]
assert list(loan1.columns) == [c for c in loan0.columns if c != "lead_id"]
pd.testing.assert_frame_equal(loan0.drop(columns="lead_id"), loan1, check_dtype=False)
pd.testing.assert_frame_equal(lead0.drop(columns="Customer_id"), lead1.drop(columns="Customer_id"), check_dtype=False)
for k in before:
    if k not in ("fact_loan", "fact_lead"):
        pd.testing.assert_frame_equal(before[k], after[k], check_dtype=False)
fl = loan1[loan1["product_name"] == "First Loan"]
lead_by_cust = lead1.dropna(subset=["Customer_id"]).groupby("Customer_id").size()
assert fl["Customer_id"].isin(lead_by_cust.index).all(), "co khoan First Loan khong noi duoc voi lead qua Customer_id"
assert (lead_by_cust.reindex(fl["Customer_id"]).to_numpy() == 1).all(), "khach First Loan co >1 lead da chuyen doi"
inst = after["fact_app_install"].merge(lead1[["lead_id", "Customer_id"]].dropna(), on="lead_id")
linked = fl.merge(inst[["Customer_id", "phone_verified_time"]].dropna(), on="Customer_id")
assert len(linked) == len(fl) and (linked["create_at"] > linked["phone_verified_time"]).all()
print(f"Da bo cot lead_id khoi fact_loan ({len(loan1)} dong, {len(loan1.columns)} cot). Bo sung Customer_id cho {n_filled} lead First Loan da chuyen doi.")
print(f"Noi Xac thuc SDT -> Dang ky qua Customer_id: {len(linked)}/{len(fl)} khoan First Loan, ket qua dung thu tu thoi gian. Cac sheet khac khong doi.")
